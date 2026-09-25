"""Synthetic-only raster and package security checks for the isolated worker."""

import io
import uuid
import warnings
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from docx import Document
from PIL import Image, UnidentifiedImageError
from pypdf import PdfWriter
from pypdf.generic import (
    DecodedStreamObject,
    DictionaryObject,
    NameObject,
    NumberObject,
    StreamObject,
)

from meyar.photo import policy
from meyar.photo.worker import UnsafeImage, _sanitize, extract
from meyar.storage.photo import LocalPhotoStorage


def _raster(
    size: tuple[int, int] = (320, 420), *, color: str = "#4585a4", fmt: str = "PNG"
) -> bytes:
    output = io.BytesIO()
    Image.new("RGB", size, color).save(output, format=fmt)
    return output.getvalue()


def _docx(*rasters: bytes) -> bytes:
    document = Document()
    document.add_paragraph("SYNTHETIC DOCUMENT")
    for raster in rasters:
        document.add_picture(io.BytesIO(raster))
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def _pdf(*rasters: bytes, page_size: tuple[int, int] = (612, 792)) -> bytes:
    writer = PdfWriter()
    page = writer.add_blank_page(width=page_size[0], height=page_size[1])
    xobjects = DictionaryObject()
    commands = []
    for index, raster in enumerate(rasters):
        with Image.open(io.BytesIO(raster)) as image:
            jpeg = io.BytesIO()
            image.convert("RGB").save(jpeg, format="JPEG")
            width, height = image.size
        stream = StreamObject()
        stream._data = jpeg.getvalue()
        stream.update(
            {
                NameObject("/Type"): NameObject("/XObject"),
                NameObject("/Subtype"): NameObject("/Image"),
                NameObject("/Width"): NumberObject(width),
                NameObject("/Height"): NumberObject(height),
                NameObject("/BitsPerComponent"): NumberObject(8),
                NameObject("/ColorSpace"): NameObject("/DeviceRGB"),
                NameObject("/Filter"): NameObject("/DCTDecode"),
            }
        )
        name = NameObject(f"/Im{index}")
        xobjects[name] = writer._add_object(stream)
        commands.append(f"q 120 0 0 170 30 {390 - index * 180} cm {name} Do Q")
    page[NameObject("/Resources")][NameObject("/XObject")] = xobjects
    content = DecodedStreamObject()
    content.set_data("\n".join(commands).encode("ascii"))
    page[NameObject("/Contents")] = writer._add_object(content)
    output = io.BytesIO()
    writer.write(output)
    return output.getvalue()


def _repack(data: bytes, extra: list[tuple[str, bytes]]) -> bytes:
    output = io.BytesIO()
    with ZipFile(io.BytesIO(data)) as source, ZipFile(output, "w", ZIP_DEFLATED) as dest:
        for name in source.namelist():
            dest.writestr(name, source.read(name))
        for name, body in extra:
            dest.writestr(name, body)
    return output.getvalue()


def _replace_media(data: bytes, replacement: bytes) -> bytes:
    output = io.BytesIO()
    with ZipFile(io.BytesIO(data)) as source, ZipFile(output, "w", ZIP_DEFLATED) as dest:
        for name in source.namelist():
            dest.writestr(
                name, replacement if name.startswith("word/media/") else source.read(name)
            )
    return output.getvalue()


def test_docx_one_portrait_and_ambiguous_images() -> None:
    one = extract(_docx(_raster()), "DOCX")
    assert one["status"] == "AVAILABLE"
    assert one["source_kind"] == "DOCX"
    assert one["source_locator"].startswith("word/media/")
    assert one["source_page"] is None
    assert (
        extract(_docx(_raster(color="#4585a4"), _raster(color="#a45264")), "DOCX")["status"]
        == "AMBIGUOUS"
    )


def test_pdf_portrait_logo_reuse_and_scan() -> None:
    portrait = _raster()
    result = extract(_pdf(_raster((48, 48)), portrait), "PDF")
    assert result["status"] == "AVAILABLE"
    assert result["source_page"] == 1
    assert "xobject-" in result["source_locator"]
    assert "object-" in result["source_locator"]
    assert extract(_pdf(portrait, portrait), "PDF")["status"] == "AVAILABLE"
    assert extract(_pdf(_raster((1224, 1584))), "PDF")["status"] == "UNUSABLE"


def test_docx_package_traversal_duplicate_external_and_limits() -> None:
    baseline = _docx(_raster())
    assert extract(_repack(baseline, [("../escape", b"x")]), "DOCX")["status"] == "UNUSABLE"
    with warnings.catch_warnings():
        warnings.simplefilter("ignore", UserWarning)
        duplicated = _repack(baseline, [("word/document.xml", b"<bad/>")])
    assert extract(duplicated, "DOCX")["status"] == "UNUSABLE"
    assert (
        extract(
            _repack(
                baseline, [("word/media/huge.png", b"x" * (policy.MAX_SOURCE_IMAGE_BYTES + 1))]
            ),
            "DOCX",
        )["status"]
        == "UNUSABLE"
    )
    assert extract(_docx(), "DOCX")["status"] == "NO_PHOTO"
    with ZipFile(io.BytesIO(baseline)) as source:
        rels = source.read("word/_rels/document.xml.rels")
    external = rels.replace(
        b'Target="media/image1.png"',
        b'Target="https://invalid.example/image.png" TargetMode="External"',
    )
    output = io.BytesIO()
    with ZipFile(io.BytesIO(baseline)) as source, ZipFile(output, "w", ZIP_DEFLATED) as dest:
        for name in source.namelist():
            dest.writestr(
                name, external if name == "word/_rels/document.xml.rels" else source.read(name)
            )
    assert extract(output.getvalue(), "DOCX")["status"] == "NO_PHOTO"


def test_sanitization_is_deterministic_and_strips_metadata() -> None:
    image = Image.new("RGB", (320, 420), "#4585a4")
    exif = Image.Exif()
    exif[270] = "synthetic private description"
    source = io.BytesIO()
    image.save(source, format="JPEG", exif=exif)
    first, width, height = _sanitize(source.getvalue())
    second, _, _ = _sanitize(source.getvalue())
    assert first == second
    assert (width, height) == (320, 420)
    with Image.open(io.BytesIO(first)) as derived:
        assert derived.format == "JPEG"
        assert not derived.getexif()
        assert "icc_profile" not in derived.info
    for mode in ("RGBA", "CMYK"):
        source = io.BytesIO()
        Image.new(mode, (320, 420)).save(source, format="PNG" if mode == "RGBA" else "JPEG")
        assert _sanitize(source.getvalue())[0].startswith(b"\xff\xd8")


def test_unsupported_malformed_and_boundary_rasters_fail_safely() -> None:
    assert extract(_docx(_raster((24, 24))), "DOCX")["status"] == "NO_PHOTO"
    assert extract(_replace_media(_docx(_raster()), b"bad raster"), "DOCX")["status"] == "UNUSABLE"
    assert extract(_docx(_raster(fmt="GIF")), "DOCX")["status"] == "UNUSABLE"
    assert extract(_docx(_raster((policy.MAX_DIMENSION + 1, 420))), "DOCX")["status"] in (
        "NO_PHOTO",
        "UNUSABLE",
    )
    with pytest.raises(UnidentifiedImageError):
        _sanitize(b"bad raster")


def test_image_count_and_raster_byte_pixel_limits(monkeypatch) -> None:
    many = [_raster(color=f"#{index:02x}4567") for index in range(policy.MAX_IMAGE_OBJECTS + 1)]
    assert extract(_docx(*many), "DOCX")["status"] == "UNUSABLE"
    with pytest.raises(UnsafeImage):
        _sanitize(_raster((policy.MAX_DIMENSION + 1, 420)))
    monkeypatch.setattr(policy, "MAX_SOURCE_IMAGE_BYTES", 20)
    with pytest.raises(UnsafeImage):
        _sanitize(_raster())
    monkeypatch.setattr(policy, "MAX_SOURCE_IMAGE_BYTES", 3 * 1024 * 1024)
    monkeypatch.setattr(policy, "MAX_DERIVED_BYTES", 20)
    with pytest.raises(UnsafeImage):
        _sanitize(_raster())


def test_malformed_pdf_image_fails_safely() -> None:
    pdf = _pdf(_raster())
    corrupted = pdf.replace(b"\xff\xd8", b"zz", 1)
    assert corrupted != pdf
    assert extract(corrupted, "PDF")["status"] in ("UNUSABLE", "EXTRACTION_FAILED")


async def test_isolated_worker_timeout_and_failure(tmp_path: Path, monkeypatch) -> None:
    from meyar.services import candidate_photo_service

    sleeper = tmp_path / "sleep-worker"
    sleeper.write_text("#!/bin/sh\nexec sleep 5\n")
    sleeper.chmod(0o700)
    monkeypatch.setattr(candidate_photo_service.sys, "executable", str(sleeper))
    monkeypatch.setattr(policy, "WORKER_TIMEOUT_SECONDS", 0.05)
    timed_out = await candidate_photo_service._extract_isolated(_docx(), "DOCX")
    assert timed_out == {"status": "EXTRACTION_FAILED", "reason_code": "WORKER_TIMEOUT"}
    sleeper.write_text("#!/bin/sh\nexit 4\n")
    failed = await candidate_photo_service._extract_isolated(_docx(), "DOCX")
    assert failed == {"status": "EXTRACTION_FAILED", "reason_code": "WORKER_FAILED"}


async def test_derived_storage_enforces_tenant_namespace(tmp_path: Path) -> None:
    storage = LocalPhotoStorage(str(tmp_path))
    owner, foreign = uuid.uuid4(), uuid.uuid4()
    key = await storage.save(tenant_id=owner, content=b"synthetic")
    assert await storage.read(tenant_id=owner, storage_key=key) == b"synthetic"
    with pytest.raises(ValueError):
        await storage.read(tenant_id=foreign, storage_key=key)
    with pytest.raises(ValueError):
        await storage.delete(tenant_id=foreign, storage_key=key)
    assert await storage.read(tenant_id=owner, storage_key=key) == b"synthetic"
