"""Independent review blockers: synthetic omitted context and reference identity."""

import copy
import json
from types import SimpleNamespace
from unittest.mock import AsyncMock

import docx
import pytest
from docx.oxml import OxmlElement
from test_docx_tables import canonical, claim, parse, ref, row_document, textbox

from meyar.extraction.evidence import (
    EvidenceValidationError,
    verify_extraction_evidence,
    verify_identity_evidence,
)
from meyar.extraction.view import build_identity_document_view, build_professional_document_view
from meyar.ingestion import parser_supervisor
from meyar.ingestion.parser import ParseError, ParseFailureCode
from meyar.ingestion.parser_policy import OutputLimits
from meyar.schemas.candidate_identity import CandidateIdentityExtraction, IdentityFieldItem
from meyar.services.identity_authority import EMPTY_IDENTITY_VALUES, identity_values_from_version
from meyar.services.profile_authority import ProfileAuthorityError, _verify_against_canonical


def omitted(cell, kind, text="No experience with"):
    if kind == "textbox":
        textbox(cell, text)
        return
    wrapper = OxmlElement(f"w:{kind}")
    parent = wrapper
    if kind == "sdt":
        parent = OxmlElement("w:sdtContent")
        wrapper.append(parent)
    p, r, t = (
        OxmlElement("w:p"),
        OxmlElement("w:r"),
        OxmlElement("w:delText" if kind == "del" else "w:t"),
    )
    t.text = text
    r.append(t)
    p.append(r)
    parent.append(p)
    cell._tc.append(wrapper)


@pytest.mark.parametrize("kind", ["sdt", "textbox", "ins", "del", "moveFrom", "moveTo"])
def test_omitted_row_rejects_positive_and_current_profile_authority(kind, caplog):
    document = row_document("", "Python")
    omitted(document.tables[0].cell(0, 0), kind)
    result = parse(document)
    stored = canonical(result)
    view = build_professional_document_view(stored)
    index = next(b.block_index for b in view.blocks if b.text == "Python")
    proposal = claim(index)
    with pytest.raises(EvidenceValidationError):
        verify_extraction_evidence(view, proposal)
    stored.tenant_id = stored.id
    stored.candidate_document_id = stored.id
    version = SimpleNamespace(
        canonical_document_id=stored.id, tenant_id=stored.id, candidate_document_id=stored.id
    )
    with pytest.raises(ProfileAuthorityError):
        _verify_against_canonical(version, proposal, stored)
    assert "No experience with" not in result.model_dump_json()
    assert "No experience with" not in view.model_dump_json() + caplog.text
    assert "row_context_complete" not in view.model_dump_json()


@pytest.mark.parametrize(
    "label,field,value",
    [
        ("Reference", "full_name", "Jane Doe"),
        ("Reference email", "email", "ref@example.com"),
        ("Emergency contact", "full_name", "Jane Doe"),
        ("Referee", "phone", "+994 50 123 45 67"),
        ("Recommender", "email", "ref@example.com"),
    ],
)
async def test_reference_identity_rejected_on_extraction_and_persisted_rebuild(label, field, value):
    stored = canonical(parse(row_document(label, value)))
    identity = CandidateIdentityExtraction(
        **{field: IdentityFieldItem(value=value, evidence=[ref(1, value)])}
    )
    with pytest.raises(EvidenceValidationError):
        verify_identity_evidence(build_identity_document_view(stored), identity)
    version = SimpleNamespace(
        status="COMPLETED",
        identity_content=identity.model_dump(mode="json"),
        canonical_document_id=stored.id,
        tenant_id=stored.id,
        candidate_document_id=stored.id,
    )
    original = copy.deepcopy(version.identity_content)
    assert (
        await identity_values_from_version(
            SimpleNamespace(scalar=AsyncMock(return_value=stored)), version=version
        )
        == EMPTY_IDENTITY_VALUES
    )
    assert version.identity_content == original


@pytest.mark.parametrize("kind", ["sdt", "textbox", "ins", "del"])
def test_incomplete_context_is_immediate_row_local_with_body_unchanged(kind):
    document = row_document("Python", "")
    other = document.tables[0].add_row()
    omitted(other.cells[0], kind)
    other.cells[1].text = "SQL"
    document.add_paragraph("Python")
    stored = canonical(parse(document))
    view = build_professional_document_view(stored)
    verify_extraction_evidence(view, claim(0))
    body = next(b for b in view.blocks if b.text == "Python" and b.block_index != 0)
    verify_extraction_evidence(view, claim(body.block_index))
    sql = next(b for b in view.blocks if b.text == "SQL")
    with pytest.raises(EvidenceValidationError):
        verify_extraction_evidence(view, claim(sql.block_index, "SQL", "SQL"))


def test_nested_row_is_separate_and_complete_hyperlink_text_stays_supported():
    document = row_document("Python", "")
    cell = document.tables[0].cell(0, 1)
    nested = cell.add_table(rows=1, cols=2)
    omitted(nested.cell(0, 0), "sdt")
    nested.cell(0, 1).text = "SQL"
    link, run, text = OxmlElement("w:hyperlink"), OxmlElement("w:r"), OxmlElement("w:t")
    text.text = "label"
    run.append(text)
    link.append(run)
    cell.paragraphs[0]._p.append(link)
    view = build_professional_document_view(canonical(parse(document)))
    verify_extraction_evidence(view, claim(0))
    sql = next(b for b in view.blocks if b.text == "SQL")
    with pytest.raises(EvidenceValidationError):
        verify_extraction_evidence(view, claim(sql.block_index, "SQL", "SQL"))


@pytest.mark.parametrize(
    "field,value",
    [
        ("full_name", "Jane Doe"),
        ("email", "candidate@example.com"),
        ("phone", "+994 50 123 45 67"),
    ],
)
def test_complete_table_and_body_identity_preserved_and_incomplete_table_rejected(field, value):
    identity = CandidateIdentityExtraction(
        **{field: IdentityFieldItem(value=value, evidence=[ref(0, value)])}
    )
    document = row_document(value, "")
    verify_identity_evidence(build_identity_document_view(canonical(parse(document))), identity)
    omitted(document.tables[0].cell(0, 1), "sdt", "Reference")
    with pytest.raises(EvidenceValidationError):
        verify_identity_evidence(build_identity_document_view(canonical(parse(document))), identity)
    body = docx.Document()
    body.add_paragraph(value)
    body.add_paragraph("Reference")
    verify_identity_evidence(build_identity_document_view(canonical(parse(body))), identity)


@pytest.mark.parametrize(
    "field,value,label,quote",
    [
        ("full_name", "Jane Doe", "Candidate name Jane Doe", "Jane"),
        ("email", "candidate@example.com", "Candidate email", "example.com"),
        ("phone", "994501234567", "Phone", "994501234567"),
    ],
)
def test_positive_identity_sibling_cannot_supply_missing_same_block_support(
    field, value, label, quote
):
    view = build_identity_document_view(canonical(parse(row_document(label, quote))))
    identity = CandidateIdentityExtraction(
        **{field: IdentityFieldItem(value=value, evidence=[ref(1, quote)])}
    )
    with pytest.raises(EvidenceValidationError):
        verify_identity_evidence(view, identity)


@pytest.mark.parametrize(
    "fault", ["missing", "null", "string", "int", "inconsistent", "no_warning"]
)
def test_row_completeness_strict_worker_parent_contract(fault):
    document = row_document("Python", "label")
    omitted(document.tables[0].cell(0, 1), "sdt")
    payload = parse(document).model_dump(mode="json")
    blocks = payload["content"]["pages"][0]["blocks"]
    source = blocks[0]["source"]
    if fault == "missing":
        del source["row_context_complete"]
    elif fault == "no_warning":
        payload["content"]["warnings"] = []
    elif fault == "inconsistent":
        source["row_context_complete"] = True
    else:
        source["row_context_complete"] = {"null": None, "string": "false", "int": 0}[fault]
    with pytest.raises(ParseError) as caught:
        parser_supervisor.decode_result(json.dumps(payload).encode(), "DOCX", OutputLimits())
    assert caught.value.code == ParseFailureCode.INVALID_PARSER_OUTPUT


def test_historical_table_missing_completeness_is_unknown_without_rewrite():
    stored = canonical(parse(row_document("Python")))
    del stored.content["pages"][0]["blocks"][0]["source"]["row_context_complete"]
    before = copy.deepcopy(stored.content)
    with pytest.raises(EvidenceValidationError):
        verify_extraction_evidence(build_professional_document_view(stored), claim(0))
    assert stored.content == before


@pytest.mark.parametrize("kind", ["sdt", "ins", "del"])
def test_run_revision_and_content_control_omissions_are_not_extracted(kind):
    document = row_document("Python")
    cell = document.tables[0].cell(0, 0)
    omitted(cell, kind)
    # Put the wrapper inside the ordinary paragraph as an inline structure.
    cell.paragraphs[0]._p.append(cell._tc[-1])
    result = parse(document)
    assert result.content.warnings == ["DOCX_TABLE_TEXT_OMITTED"]
    with pytest.raises(EvidenceValidationError):
        verify_extraction_evidence(build_professional_document_view(canonical(result)), claim(0))
