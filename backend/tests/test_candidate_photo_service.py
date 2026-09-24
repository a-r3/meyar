"""Synthetic photo provenance and presentation authority checks."""

import hashlib
import io
import uuid
from datetime import date
from pathlib import Path

import pytest
from docx import Document
from httpx import AsyncClient
from sqlalchemy import select
from sqlalchemy.exc import IntegrityError
from sqlalchemy.ext.asyncio import AsyncSession

from meyar.config import Settings, get_settings
from meyar.ingestion.parsers.local_text_parser import LocalTextParser
from meyar.main import app
from meyar.models.candidate import Candidate
from meyar.models.candidate_document import CandidateDocument
from meyar.models.candidate_photo_version import CandidatePhotoVersion
from meyar.models.canonical_document import CanonicalDocument
from meyar.models.folder_indexed_file import FolderIndexedFile
from meyar.models.folder_source import FolderSource
from meyar.services.candidate_document_repo import (
    create_candidate_document,
    get_latest_canonical_document,
)
from meyar.services.candidate_document_service import ingest_candidate_document
from meyar.services.candidate_identity_repo import (
    create_identity_version,
    get_current_identity_version,
)
from meyar.services.candidate_photo_service import (
    PLACEHOLDER_JPEG,
    current_presentable_photo,
    process_photo_for_document,
)
from meyar.services.candidate_profile_repo import (
    create_profile_version,
    get_current_profile_version,
)
from meyar.services.candidate_repo import create_candidate
from meyar.services.demo_seed_service import (
    _build_docx,
    _demo_candidates,
    _synthetic_portrait,
    seed_demo,
)
from meyar.services.tenant_repo import create_tenant
from meyar.storage.local import LocalFilesystemStorage
from meyar.storage.photo import LocalPhotoStorage

MAX_BYTES = 10 * 1024 * 1024


async def test_photo_save_failure_cleans_temporary_asset(tmp_path: Path, monkeypatch) -> None:
    from meyar.storage import photo as photo_module

    tenant_id = uuid.uuid4()
    photos = LocalPhotoStorage(str(tmp_path / "storage"))

    def fail_replace(*args, **kwargs):
        raise OSError("synthetic atomic photo write failure")

    monkeypatch.setattr(photo_module.os, "replace", fail_replace)
    with pytest.raises(OSError):
        await photos.save(tenant_id=tenant_id, content=b"synthetic derived bytes")
    assert not list((tmp_path / "storage" / "photo" / tenant_id.hex).glob("*"))


async def _seed(db: AsyncSession, tmp_path: Path):
    storage = LocalFilesystemStorage(root=str(tmp_path / "storage"))
    summary = await seed_demo(
        db,
        storage,
        LocalTextParser(),
        max_bytes=MAX_BYTES,
        max_profile_input_chars=20000,
        max_identity_input_chars=20000,
        max_embedding_input_chars=20000,
        evaluation_as_of_date=date(2026, 1, 1),
    )
    await db.commit()
    document = await db.scalar(
        select(CandidateDocument).where(
            CandidateDocument.tenant_id == summary.tenant_id,
            CandidateDocument.original_filename == _demo_candidates()[0].doc_filename,
        )
    )
    assert document is not None
    return summary, storage, document


async def _changed(
    db: AsyncSession,
    storage: LocalFilesystemStorage,
    old: CandidateDocument,
    portrait_index: int | None,
    data: bytes | None = None,
) -> CandidateDocument:
    document = await ingest_candidate_document(
        db,
        storage,
        LocalTextParser(),
        tenant_id=old.tenant_id,
        candidate_id=old.candidate_id,
        filename="synthetic-changed.docx",
        content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        data=data if data is not None else _build_docx(_demo_candidates()[0].lines, portrait_index),
        max_bytes=MAX_BYTES,
    )
    await db.commit()
    prior_identity = await get_current_identity_version(
        db, tenant_id=old.tenant_id, candidate_id=old.candidate_id
    )
    canonical = await get_latest_canonical_document(
        db, tenant_id=old.tenant_id, candidate_document_id=document.id
    )
    assert prior_identity is not None and canonical is not None
    await create_identity_version(
        db,
        tenant_id=old.tenant_id,
        candidate_id=old.candidate_id,
        candidate_document_id=document.id,
        canonical_document_id=canonical.id,
        source_sha256=document.sha256_hash,
        schema_version=prior_identity.schema_version,
        prompt_version=prior_identity.prompt_version,
        model_provider=prior_identity.model_provider,
        model_name=prior_identity.model_name,
        status="COMPLETED",
        identity_content=prior_identity.identity_content,
    )
    await db.commit()
    return document


async def test_freshness_reuse_and_no_photo(db_session: AsyncSession, tmp_path: Path) -> None:
    _summary, storage, old = await _seed(db_session, tmp_path)
    photos = LocalPhotoStorage(str(tmp_path / "storage"))

    async def process(document: CandidateDocument):
        return await process_photo_for_document(
            db_session,
            storage,
            photos,
            tenant_id=document.tenant_id,
            candidate_id=document.candidate_id,
            document_id=document.id,
        )

    first = await process(old)
    assert first is not None and first.status == "AVAILABLE"
    assert first.source_page is None and first.source_locator.startswith("word/media/")
    assert (await process(old)).id == first.id
    assert (
        await current_presentable_photo(
            db_session, tenant_id=old.tenant_id, candidate_id=old.candidate_id
        )
    ).id == first.id

    changed = await _changed(db_session, storage, old, portrait_index=1)
    assert (
        await current_presentable_photo(
            db_session, tenant_id=old.tenant_id, candidate_id=old.candidate_id
        )
        is None
    )
    second = await process(changed)
    assert second is not None and second.status == "AVAILABLE"
    assert second.version_number == first.version_number + 1
    assert second.derived_sha256 != first.derived_sha256
    assert (
        await current_presentable_photo(
            db_session, tenant_id=old.tenant_id, candidate_id=old.candidate_id
        )
    ).id == second.id

    no_photo = await _changed(db_session, storage, old, portrait_index=None)
    third = await process(no_photo)
    assert third is not None and third.status == "NO_PHOTO"
    assert (
        await current_presentable_photo(
            db_session, tenant_id=old.tenant_id, candidate_id=old.candidate_id
        )
        is None
    )


async def test_photo_constraints_reject_mismatches(
    db_session: AsyncSession, tmp_path: Path
) -> None:
    _summary, _storage, document = await _seed(db_session, tmp_path)
    tenant_id, candidate_id, document_id = (document.tenant_id, document.candidate_id, document.id)
    bad = CandidatePhotoVersion(
        tenant_id=tenant_id,
        candidate_id=candidate_id,
        candidate_document_id=document_id,
        version_number=1,
        status="NO_PHOTO",
        extractor_version="photo-v1",
        derived_storage_key="photo/invalid",
    )
    db_session.add(bad)
    with pytest.raises(IntegrityError):
        await db_session.commit()
    await db_session.rollback()
    other = await create_candidate(db_session, tenant_id=tenant_id)
    await db_session.commit()
    bad = CandidatePhotoVersion(
        tenant_id=tenant_id,
        candidate_id=other.id,
        candidate_document_id=document_id,
        version_number=1,
        status="NO_PHOTO",
        extractor_version="photo-v1",
    )
    db_session.add(bad)
    with pytest.raises(IntegrityError):
        await db_session.commit()
    await db_session.rollback()


async def test_photo_route_authorization_and_integrity(
    client: AsyncClient, db_session: AsyncSession, tmp_path: Path
) -> None:
    app.dependency_overrides[get_settings] = lambda: Settings(ui_cookie_secure=False)
    summary, storage, document = await _seed(db_session, tmp_path)
    photos = LocalPhotoStorage(str(tmp_path / "storage"))
    row = await process_photo_for_document(
        db_session,
        storage,
        photos,
        tenant_id=document.tenant_id,
        candidate_id=document.candidate_id,
        document_id=document.id,
    )
    assert row is not None and row.status == "AVAILABLE"
    url = f"/ui/candidates/{document.candidate_id}/photo"
    assert (await client.get(url, follow_redirects=False)).status_code in (302, 303)
    login = await client.post(
        "/ui/login",
        data={"username": summary.human_username, "password": summary.human_temp_password},
        follow_redirects=False,
    )
    assert login.status_code == 303
    response = await client.get(url)
    assert response.status_code == 200
    assert response.headers["content-type"].startswith("image/jpeg")
    assert response.headers["cache-control"] == "no-store"
    assert response.headers["x-content-type-options"] == "nosniff"
    assert hashlib.sha256(response.content).hexdigest() == row.derived_sha256
    assert row.derived_storage_key is not None
    assert row.derived_storage_key not in str(response.headers)
    assert row.derived_storage_key not in (await client.get("/ui/library")).text
    foreign = await create_tenant(db_session, name="Synthetic foreign")
    foreign_candidate = await create_candidate(db_session, tenant_id=foreign.id)
    await db_session.commit()
    assert (await client.get(f"/ui/candidates/{foreign_candidate.id}/photo")).status_code == 404
    await photos.delete(tenant_id=document.tenant_id, storage_key=row.derived_storage_key)
    missing = await client.get(url)
    assert missing.status_code == 200 and missing.content == PLACEHOLDER_JPEG


async def test_failed_photo_upload_keeps_durable_document(
    client: AsyncClient, db_session: AsyncSession, tmp_path: Path, monkeypatch
) -> None:
    from meyar.services import candidate_photo_service

    _summary, _storage, original = await _seed(db_session, tmp_path)

    async def fail_worker(_data: bytes, _kind: str) -> dict:
        raise RuntimeError("synthetic worker failure")

    monkeypatch.setattr(candidate_photo_service, "_extract_isolated", fail_worker)
    auth = {"Authorization": f"Bearer {_summary.api_key_plaintext}"}
    uploaded = await client.post(
        f"/api/v1/candidates/{original.candidate_id}/documents",
        headers=auth,
        files={
            "file": (
                "synthetic-changed.docx",
                _build_docx(_demo_candidates()[0].lines, 1),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )
    assert uploaded.status_code == 201
    document_id = uuid.UUID(uploaded.json()["id"])
    assert uploaded.json()["parser_status"] == "PARSED"
    row = await db_session.scalar(
        select(CandidatePhotoVersion).where(
            CandidatePhotoVersion.candidate_document_id == document_id
        )
    )
    assert row is not None and row.status == "EXTRACTION_FAILED"
    document = await db_session.get(CandidateDocument, document_id)
    assert document is not None
    assert await LocalFilesystemStorage(str(tmp_path / "storage")).read(
        storage_key=document.storage_key
    )


@pytest.mark.parametrize(
    "fault",
    [
        "document_lookup", "existing_photo_lookup", "original_read",
        "worker", "photo_save", "photo_insert", "photo_commit",
    ],
)
async def test_upload_photo_faults_never_undo_durable_ingestion(
    client: AsyncClient, db_session: AsyncSession, tmp_path: Path, monkeypatch, fault: str
) -> None:
    from meyar.services import candidate_photo_service

    summary, storage, original = await _seed(db_session, tmp_path)

    async def fail(*args, **kwargs):
        raise OSError(f"synthetic {fault} failure")

    if fault == "document_lookup":
        monkeypatch.setattr(candidate_photo_service, "get_candidate_document", fail)
    elif fault == "existing_photo_lookup":
        monkeypatch.setattr(candidate_photo_service, "get_photo_for_document", fail)
    elif fault == "original_read":
        monkeypatch.setattr(storage, "read", fail)
        from meyar.storage.dependency import get_document_storage
        app.dependency_overrides[get_document_storage] = lambda: storage
    elif fault == "worker":
        monkeypatch.setattr(candidate_photo_service, "_extract_isolated", fail)
    elif fault == "photo_save":
        monkeypatch.setattr(LocalPhotoStorage, "save", fail)
    elif fault == "photo_insert":
        monkeypatch.setattr(candidate_photo_service, "create_photo_version", fail)
    else:
        original_worker = candidate_photo_service._extract_isolated

        async def arm_photo_commit_failure(data: bytes, kind: str) -> dict:
            outcome = await original_worker(data, kind)
            original_commit = db_session.commit
            fail_once = True

            async def fail_photo_commit():
                nonlocal fail_once
                if fail_once:
                    fail_once = False
                    raise OSError("synthetic photo commit failure")
                await original_commit()

            monkeypatch.setattr(db_session, "commit", fail_photo_commit)
            return outcome

        monkeypatch.setattr(candidate_photo_service, "_extract_isolated", arm_photo_commit_failure)

    response = await client.post(
        f"/api/v1/candidates/{original.candidate_id}/documents",
        headers={"Authorization": f"Bearer {summary.api_key_plaintext}"},
        files={"file": (
            "synthetic-changed.docx", _build_docx(_demo_candidates()[0].lines, 1),
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )},
    )
    assert response.status_code == 201
    document_id = uuid.UUID(response.json()["id"])
    assert response.json()["parser_status"] == "PARSED"
    document = await db_session.get(CandidateDocument, document_id)
    assert document is not None
    canonical = await db_session.scalar(
        select(CanonicalDocument).where(CanonicalDocument.candidate_document_id == document_id)
    )
    assert canonical is not None
    assert await LocalFilesystemStorage(str(tmp_path / "storage")).read(
        storage_key=document.storage_key
    )
    assert not list((tmp_path / "storage" / "photo" / document.tenant_id.hex).glob("*"))


async def test_candidate_hard_delete_removes_derived_asset(
    client: AsyncClient, db_session: AsyncSession, tmp_path: Path
) -> None:
    summary, storage, document = await _seed(db_session, tmp_path)
    photos = LocalPhotoStorage(str(tmp_path / "storage"))
    row = await process_photo_for_document(
        db_session,
        storage,
        photos,
        tenant_id=document.tenant_id,
        candidate_id=document.candidate_id,
        document_id=document.id,
    )
    assert row is not None and row.derived_storage_key is not None
    photo_key, original_key = row.derived_storage_key, document.storage_key
    photo_id, candidate_id = row.id, document.candidate_id
    assert await photos.read(tenant_id=document.tenant_id, storage_key=photo_key)
    response = await client.delete(
        f"/api/v1/candidates/{document.candidate_id}",
        headers={"Authorization": f"Bearer {summary.api_key_plaintext}"},
    )
    assert response.status_code == 204
    assert await db_session.get(Candidate, candidate_id) is None
    assert await db_session.scalar(
        select(CandidatePhotoVersion).where(CandidatePhotoVersion.id == photo_id)
    ) is None
    with pytest.raises(FileNotFoundError):
        await photos.read(tenant_id=document.tenant_id, storage_key=photo_key)
    with pytest.raises(FileNotFoundError):
        await storage.read(storage_key=original_key)


@pytest.mark.parametrize("fault", ["original_delete", "photo_delete"])
async def test_failed_hard_delete_preserves_valid_photo_and_retry(
    client: AsyncClient, db_session: AsyncSession, tmp_path: Path, monkeypatch, fault: str
) -> None:
    app.dependency_overrides[get_settings] = lambda: Settings(ui_cookie_secure=False)
    summary, storage, document = await _seed(db_session, tmp_path)
    photos = LocalPhotoStorage(str(tmp_path / "storage"))
    row = await process_photo_for_document(
        db_session, storage, photos, tenant_id=document.tenant_id,
        candidate_id=document.candidate_id, document_id=document.id,
    )
    assert row is not None and row.derived_storage_key is not None
    tenant_id, candidate_id, document_key, photo_id = (
        document.tenant_id, document.candidate_id, document.storage_key, row.id
    )
    photo_key = row.derived_storage_key
    photo_bytes = await photos.read(tenant_id=tenant_id, storage_key=photo_key)
    original_bytes = await storage.read(storage_key=document_key)
    original_delete = LocalFilesystemStorage.delete
    photo_delete = LocalPhotoStorage.delete
    enabled = True

    async def fail_original(self, *, storage_key: str) -> None:
        if enabled:
            raise OSError("synthetic original delete failure")
        await original_delete(self, storage_key=storage_key)

    async def fail_photo(self, *, tenant_id: uuid.UUID, storage_key: str) -> None:
        if enabled:
            raise OSError("synthetic photo delete failure")
        await photo_delete(self, tenant_id=tenant_id, storage_key=storage_key)

    monkeypatch.setattr(
        LocalFilesystemStorage if fault == "original_delete" else LocalPhotoStorage,
        "delete", fail_original if fault == "original_delete" else fail_photo,
    )
    url = f"/api/v1/candidates/{candidate_id}"
    auth = {"Authorization": f"Bearer {summary.api_key_plaintext}"}
    with pytest.raises(OSError):
        await client.delete(url, headers=auth)
    assert await db_session.get(Candidate, candidate_id) is not None
    assert await db_session.get(CandidatePhotoVersion, photo_id) is not None
    assert await photos.read(tenant_id=tenant_id, storage_key=photo_key) == photo_bytes
    assert await storage.read(storage_key=document_key) == original_bytes
    login = await client.post(
        "/ui/login",
        data={"username": summary.human_username, "password": summary.human_temp_password},
        follow_redirects=False,
    )
    assert login.status_code == 303
    rendered = await client.get(f"/ui/candidates/{candidate_id}/photo")
    assert rendered.status_code == 200 and rendered.content == photo_bytes
    enabled = False
    assert (await client.delete(url, headers=auth)).status_code == 204
    assert await db_session.get(Candidate, candidate_id) is None
    with pytest.raises(FileNotFoundError):
        await photos.read(tenant_id=tenant_id, storage_key=photo_key)


async def test_multi_photo_partial_delete_restores_every_available_asset(
    client: AsyncClient, db_session: AsyncSession, tmp_path: Path, monkeypatch
) -> None:
    summary, storage, old = await _seed(db_session, tmp_path)
    photos = LocalPhotoStorage(str(tmp_path / "storage"))
    first = await process_photo_for_document(
        db_session, storage, photos, tenant_id=old.tenant_id,
        candidate_id=old.candidate_id, document_id=old.id,
    )
    changed = await _changed(db_session, storage, old, portrait_index=1)
    second = await process_photo_for_document(
        db_session, storage, photos, tenant_id=changed.tenant_id,
        candidate_id=changed.candidate_id, document_id=changed.id,
    )
    assert first is not None and second is not None
    assert first.derived_storage_key and second.derived_storage_key
    tenant_id, candidate_id = old.tenant_id, old.candidate_id
    expected = {
        key: await photos.read(tenant_id=tenant_id, storage_key=key)
        for key in (first.derived_storage_key, second.derived_storage_key)
    }
    original_delete = LocalPhotoStorage.delete
    attempts = 0
    enabled = True

    async def fail_second(self, *, tenant_id: uuid.UUID, storage_key: str) -> None:
        nonlocal attempts
        attempts += 1
        if enabled and attempts == 2:
            raise OSError("synthetic second photo delete failure")
        await original_delete(self, tenant_id=tenant_id, storage_key=storage_key)

    monkeypatch.setattr(LocalPhotoStorage, "delete", fail_second)
    url = f"/api/v1/candidates/{candidate_id}"
    auth = {"Authorization": f"Bearer {summary.api_key_plaintext}"}
    with pytest.raises(OSError):
        await client.delete(url, headers=auth)
    assert await db_session.get(Candidate, candidate_id) is not None
    for key, content in expected.items():
        assert await photos.read(tenant_id=tenant_id, storage_key=key) == content
    enabled = False
    assert (await client.delete(url, headers=auth)).status_code == 204
    for key in expected:
        with pytest.raises(FileNotFoundError):
            await photos.read(tenant_id=tenant_id, storage_key=key)


async def test_photo_assets_recover_when_candidate_db_commit_fails(
    client: AsyncClient, db_session: AsyncSession, tmp_path: Path, monkeypatch
) -> None:
    summary, storage, document = await _seed(db_session, tmp_path)
    photos = LocalPhotoStorage(str(tmp_path / "storage"))
    row = await process_photo_for_document(
        db_session, storage, photos, tenant_id=document.tenant_id,
        candidate_id=document.candidate_id, document_id=document.id,
    )
    assert row is not None and row.derived_storage_key is not None
    tenant_id, candidate_id, photo_key = (
        document.tenant_id, document.candidate_id, row.derived_storage_key
    )
    photo_bytes = await photos.read(tenant_id=tenant_id, storage_key=photo_key)
    original_commit = db_session.commit
    fail_once = True

    async def fail_commit():
        nonlocal fail_once
        if fail_once:
            fail_once = False
            raise OSError("synthetic candidate DB commit failure")
        await original_commit()

    monkeypatch.setattr(db_session, "commit", fail_commit)
    url = f"/api/v1/candidates/{candidate_id}"
    auth = {"Authorization": f"Bearer {summary.api_key_plaintext}"}
    with pytest.raises(OSError):
        await client.delete(url, headers=auth)
    assert await db_session.get(Candidate, candidate_id) is not None
    assert await photos.read(tenant_id=tenant_id, storage_key=photo_key) == photo_bytes
    assert (await client.delete(url, headers=auth)).status_code == 204


async def test_hard_delete_rejects_foreign_tenant_photo_key_before_any_delete(
    client: AsyncClient, db_session: AsyncSession, tmp_path: Path
) -> None:
    summary, storage, document = await _seed(db_session, tmp_path)
    photos = LocalPhotoStorage(str(tmp_path / "storage"))
    row = await process_photo_for_document(
        db_session, storage, photos, tenant_id=document.tenant_id,
        candidate_id=document.candidate_id, document_id=document.id,
    )
    assert row is not None and row.derived_storage_key is not None
    candidate_id, document_key = document.candidate_id, document.storage_key
    original = await storage.read(storage_key=document_key)
    foreign = await create_tenant(db_session, name="Foreign photo-key tenant")
    await db_session.commit()
    foreign_key = await photos.save(tenant_id=foreign.id, content=b"foreign synthetic asset")
    row.derived_storage_key = foreign_key
    await db_session.commit()
    with pytest.raises(ValueError, match="Invalid derived photo key"):
        await client.delete(
            f"/api/v1/candidates/{candidate_id}",
            headers={"Authorization": f"Bearer {summary.api_key_plaintext}"},
        )
    assert await db_session.get(Candidate, candidate_id) is not None
    assert await storage.read(storage_key=document_key) == original
    assert (
        await photos.read(tenant_id=foreign.id, storage_key=foreign_key)
        == b"foreign synthetic asset"
    )


async def test_dedup_linked_folder_paths_diverge_without_stale_photo(
    db_session: AsyncSession, tmp_path: Path
) -> None:
    _summary, storage, old = await _seed(db_session, tmp_path)
    photos = LocalPhotoStorage(str(tmp_path / "storage"))
    old_photo = await process_photo_for_document(
        db_session,
        storage,
        photos,
        tenant_id=old.tenant_id,
        candidate_id=old.candidate_id,
        document_id=old.id,
    )
    assert old_photo is not None and old_photo.status == "AVAILABLE"
    source = FolderSource(tenant_id=old.tenant_id, root_path=str(tmp_path / "synthetic-source"))
    db_session.add(source)
    await db_session.flush()
    links = [
        FolderIndexedFile(
            tenant_id=old.tenant_id,
            folder_source_id=source.id,
            relative_path=name,
            document_type="DOCX",
            byte_size=old.byte_size,
            sha256_hash=old.sha256_hash,
            index_status="INDEXED",
            candidate_id=old.candidate_id,
            candidate_document_id=old.id,
        )
        for name in ("a.docx", "b.docx")
    ]
    db_session.add_all(links)
    await db_session.commit()
    assert links[0].candidate_document_id == links[1].candidate_document_id
    assert (
        await process_photo_for_document(
            db_session,
            storage,
            photos,
            tenant_id=old.tenant_id,
            candidate_id=old.candidate_id,
            document_id=links[1].candidate_document_id,
        )
    ).id == old_photo.id

    changed = await _changed(db_session, storage, old, portrait_index=None)
    links[0].candidate_document_id = changed.id
    links[0].sha256_hash = changed.sha256_hash
    await db_session.commit()
    assert links[1].candidate_document_id == old.id
    assert (
        await current_presentable_photo(
            db_session, tenant_id=old.tenant_id, candidate_id=old.candidate_id
        )
        is None
    )
    changed_photo = await process_photo_for_document(
        db_session,
        storage,
        photos,
        tenant_id=changed.tenant_id,
        candidate_id=changed.candidate_id,
        document_id=changed.id,
    )
    assert changed_photo is not None and changed_photo.status == "NO_PHOTO"
    assert (
        await current_presentable_photo(
            db_session, tenant_id=old.tenant_id, candidate_id=old.candidate_id
        )
        is None
    )


async def test_new_ambiguous_and_failed_photo_never_reuse_older_asset(
    db_session: AsyncSession, tmp_path: Path, monkeypatch
) -> None:
    from meyar.services import candidate_photo_service

    _summary, storage, old = await _seed(db_session, tmp_path)
    photos = LocalPhotoStorage(str(tmp_path / "storage"))
    original_photo = await process_photo_for_document(
        db_session,
        storage,
        photos,
        tenant_id=old.tenant_id,
        candidate_id=old.candidate_id,
        document_id=old.id,
    )
    assert original_photo is not None and original_photo.status == "AVAILABLE"
    document = Document()
    for line in _demo_candidates()[0].lines:
        document.add_paragraph(line)
    for index in (1, 2):
        document.add_picture(io.BytesIO(_synthetic_portrait(index)))
    output = io.BytesIO()
    document.save(output)
    ambiguous = await _changed(db_session, storage, old, None, data=output.getvalue())
    ambiguous_photo = await process_photo_for_document(
        db_session,
        storage,
        photos,
        tenant_id=ambiguous.tenant_id,
        candidate_id=ambiguous.candidate_id,
        document_id=ambiguous.id,
    )
    assert ambiguous_photo is not None and ambiguous_photo.status == "AMBIGUOUS"
    assert (
        await current_presentable_photo(
            db_session, tenant_id=old.tenant_id, candidate_id=old.candidate_id
        )
        is None
    )

    failed = await _changed(db_session, storage, old, portrait_index=3)

    async def fail_worker(_data: bytes, _kind: str) -> dict:
        return {"status": "EXTRACTION_FAILED", "reason_code": "SYNTHETIC_FAILURE"}

    monkeypatch.setattr(candidate_photo_service, "_extract_isolated", fail_worker)
    failed_photo = await process_photo_for_document(
        db_session,
        storage,
        photos,
        tenant_id=failed.tenant_id,
        candidate_id=failed.candidate_id,
        document_id=failed.id,
    )
    assert failed_photo is not None and failed_photo.status == "EXTRACTION_FAILED"
    assert (
        await current_presentable_photo(
            db_session, tenant_id=old.tenant_id, candidate_id=old.candidate_id
        )
        is None
    )


async def test_failed_profile_does_not_block_valid_identity_photo_and_parse_failure_hides_old(
    db_session: AsyncSession, tmp_path: Path
) -> None:
    _summary, storage, old = await _seed(db_session, tmp_path)
    photos = LocalPhotoStorage(str(tmp_path / "storage"))
    photo = await process_photo_for_document(
        db_session,
        storage,
        photos,
        tenant_id=old.tenant_id,
        candidate_id=old.candidate_id,
        document_id=old.id,
    )
    assert photo is not None and photo.status == "AVAILABLE"
    profile = await get_current_profile_version(
        db_session, tenant_id=old.tenant_id, candidate_id=old.candidate_id
    )
    assert profile is not None
    await create_profile_version(
        db_session,
        tenant_id=old.tenant_id,
        candidate_id=old.candidate_id,
        candidate_document_id=old.id,
        canonical_document_id=profile.canonical_document_id,
        source_sha256=old.sha256_hash,
        schema_version=profile.schema_version,
        prompt_version=profile.prompt_version,
        model_provider=profile.model_provider,
        model_name=profile.model_name,
        model_metadata={},
        status="FAILED",
        error_code="SYNTHETIC_PROFILE_FAILURE",
    )
    await db_session.commit()
    assert (
        await current_presentable_photo(
            db_session, tenant_id=old.tenant_id, candidate_id=old.candidate_id
        )
    ).id == photo.id
    data = b"synthetic parse-failed document"
    key = await storage.save(tenant_id=old.tenant_id, content=data)
    failed_document = await create_candidate_document(
        db_session,
        tenant_id=old.tenant_id,
        candidate_id=old.candidate_id,
        original_filename="synthetic-failed.pdf",
        mime_type="application/pdf",
        byte_size=len(data),
        sha256_hash=hashlib.sha256(data).hexdigest(),
        storage_key=key,
    )
    failed_document.parser_status = "PARSE_FAILED"
    await db_session.commit()
    assert (
        await current_presentable_photo(
            db_session, tenant_id=old.tenant_id, candidate_id=old.candidate_id
        )
        is None
    )
