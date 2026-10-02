"""Generated/synthetic documents only: worker lifecycle and canonical safety."""

import asyncio
import errno
import io
import json
import os
import resource
import sys
import threading
from dataclasses import replace
from pathlib import Path

import docx
import pypdf
import pytest
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject, NumberObject

from meyar.ingestion import parser_policy as policy
from meyar.ingestion import parser_supervisor as supervisor
from meyar.ingestion import parser_worker as worker
from meyar.ingestion import validation
from meyar.ingestion.admission import AdmissionGate
from meyar.ingestion.parser import PARSE_FAILURE_MESSAGES, ParseError, ParseFailureCode
from meyar.ingestion.parser_policy import OutputLimits
from meyar.ingestion.parsers.local_text_parser import LocalTextParser
from meyar.ingestion.parsers.local_text_sync import parse_sync

VALID = (Path(__file__).parents[2] / "fixtures/synthetic_cvs/valid_cv.pdf").read_bytes()


def pdf(texts: list[str | None], *, image: bool = False) -> bytes:
    writer = pypdf.PdfWriter()
    for text in texts:
        page = writer.add_blank_page(width=72, height=72)
        if text:
            font = DictionaryObject(
                {
                    NameObject("/Type"): NameObject("/Font"),
                    NameObject("/Subtype"): NameObject("/Type1"),
                    NameObject("/BaseFont"): NameObject("/Helvetica"),
                }
            )
            page[NameObject("/Resources")] = DictionaryObject(
                {NameObject("/Font"): DictionaryObject({NameObject("/F1"): font})}
            )
            stream = DecodedStreamObject()
            stream.set_data(f"BT /F1 12 Tf ({text}) Tj ET".encode())
            page[NameObject("/Contents")] = writer._add_object(stream)
        if image:
            raster = DecodedStreamObject()
            raster.set_data(b"\x00\x00\x00")
            raster.update(
                {
                    NameObject("/Type"): NameObject("/XObject"),
                    NameObject("/Subtype"): NameObject("/Image"),
                    NameObject("/Width"): NumberObject(1),
                    NameObject("/Height"): NumberObject(1),
                    NameObject("/ColorSpace"): NameObject("/DeviceRGB"),
                    NameObject("/BitsPerComponent"): NumberObject(8),
                }
            )
            page[NameObject("/Resources")] = DictionaryObject(
                {
                    NameObject("/XObject"): DictionaryObject(
                        {NameObject("/Im1"): writer._add_object(raster)}
                    )
                }
            )
            stream = DecodedStreamObject()
            stream.set_data(b"q 72 0 0 72 0 0 cm /Im1 Do Q")
            page[NameObject("/Contents")] = writer._add_object(stream)
    buffer = io.BytesIO()
    writer.write(buffer)
    return buffer.getvalue()


def docx_bytes(paragraphs: list[str], *, table: bool = False) -> bytes:
    document = docx.Document()
    for text in paragraphs:
        document.add_paragraph(text)
    if table:
        document.add_table(rows=1, cols=1).cell(0, 0).text = "Synthetic Python engineer"
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


async def normal() -> None:
    result = await LocalTextParser().parse(data=VALID, document_type="PDF")
    assert result.parser_version == "1.2.0"
    assert result.content.pages[0].blocks


def assert_reaped(process: asyncio.subprocess.Process) -> None:
    assert process.returncode is not None
    with pytest.raises(ProcessLookupError):
        os.kill(process.pid, 0)
    with pytest.raises(ChildProcessError):
        os.waitpid(process.pid, os.WNOHANG)


@pytest.mark.parametrize(
    ("source", "code"),
    [
        ("import time; time.sleep(60)", ParseFailureCode.PARSER_TIMEOUT),
        (
            "import signal,time; signal.signal(signal.SIGTERM, signal.SIG_IGN); time.sleep(60)",
            ParseFailureCode.PARSER_TIMEOUT,
        ),
        ("import os; os._exit(7)", ParseFailureCode.PARSER_WORKER_FAILED),
        ('print("not JSON")', ParseFailureCode.INVALID_PARSER_OUTPUT),
        (
            'print("{\\"error\\":\\"raw /secret/path candidate text\\"}")',
            ParseFailureCode.INVALID_PARSER_OUTPUT,
        ),
        ('import sys; sys.stdout.write("x"*1000000)', ParseFailureCode.PARSER_OUTPUT_LIMIT),
        ("import os; os._exit(72)", ParseFailureCode.PARSER_RESOURCE_UNAVAILABLE),
        (
            "from meyar.ingestion.parser_worker import establish_memory_limit; "
            "establish_memory_limit();\ntry: x=bytearray(1024*1024*1024)\n"
            "except MemoryError: raise SystemExit(73)",
            ParseFailureCode.PARSER_RESOURCE_LIMIT,
        ),
        (
            "import signal,sys,time; signal.signal(signal.SIGTERM, signal.SIG_IGN); "
            'sys.stdout.write("x"*1000000); sys.stdout.flush(); time.sleep(60)',
            ParseFailureCode.PARSER_OUTPUT_LIMIT,
        ),
    ],
)
async def test_real_child_failure_reaped_and_recovery(monkeypatch, source, code):
    original = asyncio.create_subprocess_exec
    children = []

    async def fault(*args, **kwargs):
        assert kwargs["env"] == {"LANG": "C.UTF-8"}
        child = await original(sys.executable, "-c", source, **kwargs)
        children.append(child)
        return child

    with monkeypatch.context() as patch:
        patch.setattr(asyncio, "create_subprocess_exec", fault)
        with pytest.raises(ParseError) as caught:
            await supervisor.parse_isolated(
                data=VALID,
                document_type="PDF",
                limits=replace(OutputLimits(), result_bytes=512),
                seconds=0.5,
            )
        assert caught.value.code == code
        assert str(caught.value) == PARSE_FAILURE_MESSAGES[code]
    assert_reaped(children[0])
    if "SIG_IGN" in source:
        assert children[0].returncode == -9
    await normal()


async def test_startup_failure_and_recovery(monkeypatch):
    async def failure(*args, **kwargs):
        raise OSError("synthetic /private/path environment secret")

    with monkeypatch.context() as patch:
        patch.setattr(asyncio, "create_subprocess_exec", failure)
        with pytest.raises(ParseError) as caught:
            await normal()
        assert caught.value.code == ParseFailureCode.PARSER_STARTUP_FAILED
        assert "/private" not in str(caught.value)
    await normal()


async def test_timeout_waits_for_spawn_resolution_then_reaps_late_child(monkeypatch):
    original = asyncio.create_subprocess_exec
    entered = asyncio.Event()
    allow_spawn = asyncio.Event()
    children = []

    async def unresolved(*args, **kwargs):
        entered.set()
        await allow_spawn.wait()
        child = await original(sys.executable, "-c", "import time; time.sleep(60)", **kwargs)
        children.append(child)
        return child

    with monkeypatch.context() as patch:
        patch.setattr(asyncio, "create_subprocess_exec", unresolved)
        task = asyncio.create_task(
            supervisor.parse_isolated(data=VALID, document_type="PDF", seconds=0.02)
        )
        try:
            await asyncio.wait_for(entered.wait(), 3)
            await asyncio.sleep(0.08)
            # Deadline expires, but safe cleanup cannot abandon an unresolved
            # OS spawn that may still create a child. No total-time claim.
            assert not task.done() and not children
            assert supervisor.gate().active == 1
        finally:
            allow_spawn.set()
            with pytest.raises(ParseError) as caught:
                await asyncio.wait_for(task, 3)
        assert caught.value.code == ParseFailureCode.PARSER_TIMEOUT
    assert_reaped(children[0])
    assert supervisor.gate().active == 0
    await normal()


@pytest.mark.parametrize("during_startup", [False, True])
async def test_cancellation_reaps_even_during_spawn(monkeypatch, during_startup):
    original = asyncio.create_subprocess_exec
    started = asyncio.Event()
    allow_spawn = asyncio.Event()
    children = []

    async def slow(*args, **kwargs):
        child = await original(
            sys.executable,
            "-c",
            "import signal,time; signal.signal(signal.SIGTERM, signal.SIG_IGN); time.sleep(60)",
            **kwargs,
        )
        children.append(child)
        started.set()
        if during_startup:
            await allow_spawn.wait()
        return child

    with monkeypatch.context() as patch:
        patch.setattr(asyncio, "create_subprocess_exec", slow)
        task = asyncio.create_task(normal())
        await started.wait()
        if not during_startup:
            await asyncio.sleep(0.1)
        task.cancel()
        await asyncio.sleep(0)
        task.cancel()
        allow_spawn.set()
        with pytest.raises(asyncio.CancelledError):
            await asyncio.wait_for(task, 3)
    assert_reaped(children[0])
    assert supervisor.gate().active == 0
    await normal()


async def test_admission_saturation_timeout_cancel_and_recovery():
    gate = AdmissionGate(active=1, waiters=2, seconds=0.05)
    await gate.acquire()
    waiting = [asyncio.create_task(gate.acquire()) for _ in range(2)]
    await asyncio.sleep(0)
    with pytest.raises(ParseError) as caught:
        await gate.acquire()
    assert caught.value.code == ParseFailureCode.PARSER_BUSY
    waiting[0].cancel()
    with pytest.raises(asyncio.CancelledError):
        await waiting[0]
    with pytest.raises(ParseError):
        await waiting[1]
    assert not gate.waiters
    gate.release()
    async with gate.slot():
        assert gate.active == 1
    assert gate.active == 0
    await normal()


async def test_process_wide_admission_caps_actual_children(monkeypatch):
    original = asyncio.create_subprocess_exec
    gate = AdmissionGate(active=1, waiters=2, seconds=1)
    started = asyncio.Event()
    children = []

    async def slow(*args, **kwargs):
        child = await original(sys.executable, "-c", "import time; time.sleep(60)", **kwargs)
        children.append(child)
        started.set()
        return child

    with monkeypatch.context() as patch:
        patch.setattr(supervisor, "gate", lambda: gate)
        patch.setattr(asyncio, "create_subprocess_exec", slow)
        tasks = [asyncio.create_task(normal()) for _ in range(3)]
        await started.wait()
        await asyncio.sleep(0.02)
        with pytest.raises(ParseError) as caught:
            await normal()
        assert caught.value.code == ParseFailureCode.PARSER_BUSY
        assert len(children) == 1 and len(gate.waiters) == 2
        for task in tasks:
            task.cancel()
        await asyncio.gather(*tasks, return_exceptions=True)
    assert_reaped(children[0])
    await normal()


async def test_heartbeat_during_real_parsing():
    ticks = 0
    task = asyncio.create_task(
        LocalTextParser().parse(
            data=docx_bytes(["Synthetic engineer"] * 10000), document_type="DOCX"
        )
    )
    while not task.done():
        ticks += 1
        await asyncio.sleep(0.005)
    result = await task
    assert len(result.content.pages[0].blocks) == 10000
    assert ticks >= 5


@pytest.mark.parametrize(
    ("data", "kind"),
    [
        (pdf([None]), "PDF"),
        (pdf([None], image=True), "PDF"),
        (pdf([]), "PDF"),
        (docx_bytes(["  "]), "DOCX"),
    ],
)
async def test_no_extractable_text(data, kind):
    with pytest.raises(ParseError) as caught:
        await LocalTextParser().parse(data=data, document_type=kind)
    assert caught.value.code == ParseFailureCode.INSUFFICIENT_EXTRACTABLE_TEXT
    assert "scanned" not in str(caught.value).lower()
    await normal()


async def test_mixed_pdf_and_short_docx_preserve_evidence():
    result = await LocalTextParser().parse(data=pdf([None, "Hi", None]), document_type="PDF")
    assert [p.page for p in result.content.pages] == [1, 2, 3]
    assert [len(p.blocks) for p in result.content.pages] == [0, 1, 0]
    assert result.content.pages[1].blocks[0].index == 0
    result = await LocalTextParser().parse(
        data=docx_bytes(["", "Hi", "", "AZ"]), document_type="DOCX"
    )
    assert [(b.index, b.text) for b in result.content.pages[0].blocks] == [(1, "Hi"), (3, "AZ")]


async def test_exact_over_pdf_page_boundary():
    result = await LocalTextParser().parse(data=pdf(["Hi"] + [None] * 299), document_type="PDF")
    assert len(result.content.pages) == 300
    with pytest.raises(ParseError) as caught:
        await LocalTextParser().parse(data=pdf(["Hi"] + [None] * 300), document_type="PDF")
    assert caught.value.code == ParseFailureCode.PARSER_OUTPUT_LIMIT
    await normal()


@pytest.mark.parametrize(
    ("field", "texts", "exact"),
    [
        ("blocks", ["Hi", "AZ"], 2),
        ("characters", ["Python"], 6),
        ("text_bytes", ["Azərbaycan Русский 中文"], len("Azərbaycan Русский 中文".encode())),
    ],
)
async def test_exact_over_incremental_limits(field, texts, exact):
    data = docx_bytes(texts)
    limits = replace(OutputLimits(), **{field: exact})
    result = await supervisor.parse_isolated(data=data, document_type="DOCX", limits=limits)
    assert len(result.content.pages[0].blocks) == len(texts)
    with pytest.raises(ParseError) as caught:
        await supervisor.parse_isolated(
            data=data, document_type="DOCX", limits=replace(limits, **{field: exact - 1})
        )
    assert caught.value.code == ParseFailureCode.PARSER_OUTPUT_LIMIT
    # Parent independently applies the same bounds to a valid worker-shaped result.
    with pytest.raises(ParseError) as caught:
        supervisor.decode_result(
            result.model_dump_json().encode(), "DOCX", replace(limits, **{field: exact - 1})
        )
    assert caught.value.code == ParseFailureCode.PARSER_OUTPUT_LIMIT
    await normal()


async def test_exact_over_serialized_limit():
    data = docx_bytes(['A\tB " Unicode ə'])
    result = parse_sync(data, "DOCX", OutputLimits())
    encoded = worker.encode_bounded(result.model_dump(mode="json"), 10000)
    limit = replace(OutputLimits(), result_bytes=len(encoded))
    assert await supervisor.parse_isolated(data=data, document_type="DOCX", limits=limit)
    for operation in (
        lambda: worker.encode_bounded(result.model_dump(mode="json"), len(encoded) - 1),
        lambda: supervisor.decode_result(
            encoded, "DOCX", replace(limit, result_bytes=len(encoded) - 1)
        ),
    ):
        with pytest.raises(ParseError) as caught:
            operation()
        assert caught.value.code == ParseFailureCode.PARSER_OUTPUT_LIMIT
    with pytest.raises(ParseError) as caught:
        await supervisor.parse_isolated(
            data=data, document_type="DOCX", limits=replace(limit, result_bytes=len(encoded) - 1)
        )
    assert caught.value.code == ParseFailureCode.PARSER_OUTPUT_LIMIT
    await normal()


@pytest.mark.parametrize(
    "output",
    [
        b"{}",
        b"[]",
        b'{"error":"oops"}',
        b'{"error":"INVALID_DOCUMENT","extra":"secret"}',
        b'{"a":1,"a":2}',
        b"\xff",
    ],
)
def test_closed_worker_schema(output):
    with pytest.raises(ParseError) as caught:
        supervisor.decode_result(output, "PDF", OutputLimits())
    assert caught.value.code == ParseFailureCode.INVALID_PARSER_OUTPUT


@pytest.mark.parametrize(
    "fault", ["missing", "set_failed", "readback", "unenforced", "wrong_errno"]
)
@pytest.mark.parametrize("platform", ["linux", "darwin"])
def test_memory_setup_fails_closed(monkeypatch, fault, platform):
    monkeypatch.setattr(sys, "platform", platform)
    if fault == "missing":
        monkeypatch.delattr(resource, "RLIMIT_AS")
    else:

        def set_limit(*args):
            if fault == "set_failed":
                raise ValueError("synthetic failure")

        monkeypatch.setattr(resource, "setrlimit", set_limit)
        monkeypatch.setattr(
            resource,
            "getrlimit",
            lambda *_: (1, 1) if fault == "readback" else (policy.WORKER_MEMORY_BYTES,) * 2,
        )

        def map_memory(*args, **kwargs):
            if fault == "wrong_errno":
                raise OSError(errno.EINVAL, "synthetic")

            class Mapping:
                def close(self):
                    pass

            return Mapping()

        monkeypatch.setattr(worker.mmap, "mmap", map_memory)
    with pytest.raises((RuntimeError, ValueError)):
        worker.establish_memory_limit()


@pytest.mark.parametrize("platform", ["linux", "darwin"])
def test_memory_setup_requires_allocation_refusal(monkeypatch, platform):
    monkeypatch.setattr(sys, "platform", platform)
    monkeypatch.setattr(resource, "setrlimit", lambda *args: None)
    monkeypatch.setattr(resource, "getrlimit", lambda *_: (policy.WORKER_MEMORY_BYTES,) * 2)

    def refused(*args, **kwargs):
        raise OSError(errno.ENOMEM, "synthetic")

    monkeypatch.setattr(worker.mmap, "mmap", refused)
    worker.establish_memory_limit()


async def test_validation_thread_heartbeat_cancellation_and_bound(monkeypatch):
    started = threading.Event()
    release = threading.Event()
    original = validation.validate_upload

    def slow(**kwargs):
        started.set()
        release.wait(3)
        return original(**kwargs)

    gate = AdmissionGate(active=1, waiters=0)
    with monkeypatch.context() as patch:
        patch.setattr(validation, "validate_upload", slow)
        patch.setattr(validation, "gate", lambda **_: gate)
        task = asyncio.create_task(
            validation.validate_upload_async(
                filename="synthetic.docx",
                content_type="",
                data=docx_bytes(["Hi"]),
                max_bytes=100000,
            )
        )
        ticks = 0
        while not started.is_set():
            ticks += 1
            await asyncio.sleep(0.005)
        for _ in range(5):
            ticks += 1
            await asyncio.sleep(0.005)
        assert ticks >= 5
        task.cancel()
        with pytest.raises(asyncio.CancelledError):
            await task
        assert gate.active == 1
        with pytest.raises(ParseError):
            await validation.validate_upload_async(
                filename="synthetic.pdf", content_type="", data=VALID, max_bytes=100000
            )
        release.set()
        while gate.active:
            await asyncio.sleep(0.005)
    assert await validation.validate_upload_async(
        filename="synthetic.pdf", content_type="", data=VALID, max_bytes=100000
    )


@pytest.mark.parametrize(
    ("texts", "field", "bound"),
    [
        (["Hi"] * 10001, "blocks", 10000),
        (["x" * 1000001], "characters", 1000000),
    ],
)
async def test_actual_policy_exact_over(texts, field, bound):
    exact = texts[:-1] if field == "blocks" else [texts[0][:-1]]
    assert await LocalTextParser().parse(data=docx_bytes(exact), document_type="DOCX")
    with pytest.raises(ParseError) as caught:
        await LocalTextParser().parse(data=docx_bytes(texts), document_type="DOCX")
    assert caught.value.code == ParseFailureCode.PARSER_OUTPUT_LIMIT
    await normal()


@pytest.mark.parametrize("mutation", ["extra", "coerce", "page", "index", "blank", "version"])
async def test_parent_rejects_invalid_success_schema_and_recovers(mutation):
    result = parse_sync(VALID, "PDF", OutputLimits()).model_dump(mode="json")
    block = result["content"]["pages"][0]["blocks"][0]
    if mutation == "extra":
        block["unexpected"] = "synthetic private text"
    if mutation == "coerce":
        block["index"] = "0"
    if mutation == "page":
        result["content"]["pages"][0]["page"] = 2
    if mutation == "index":
        block["index"] = -1
    if mutation == "blank":
        block["text"] = " "
    if mutation == "version":
        result["parser_version"] = "1.0.0"
    with pytest.raises(ParseError) as caught:
        supervisor.decode_result(json.dumps(result).encode(), "PDF", OutputLimits())
    assert caught.value.code == ParseFailureCode.INVALID_PARSER_OUTPUT
    await normal()
