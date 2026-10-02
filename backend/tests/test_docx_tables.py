"""PR-3 source-order, omission, resource and evidence regressions; synthetic only."""

import io
import json
import uuid
from dataclasses import replace
from types import SimpleNamespace

import docx
import pytest
from docx.oxml import OxmlElement
from pydantic import ValidationError

from meyar.extraction.evidence import (
    EvidenceValidationError,
    verify_extraction_evidence,
    verify_identity_evidence,
)
from meyar.extraction.view import build_identity_document_view, build_professional_document_view
from meyar.ingestion import parser_supervisor, parser_worker
from meyar.ingestion.parser import BodySource, ParseError, ParseFailureCode, TableSource
from meyar.ingestion.parser_output import TextBudget
from meyar.ingestion.parser_policy import MAX_DOCX_TABLE_DEPTH, OutputLimits
from meyar.ingestion.parsers import docx_source
from meyar.ingestion.parsers.local_text_sync import parse_sync
from meyar.schemas.candidate_identity import CandidateIdentityExtraction, IdentityFieldItem
from meyar.schemas.candidate_profile import (
    CandidateProfileExtraction,
    EmploymentItem,
    EvidenceRef,
    SkillItem,
)
from meyar.services.profile_authority import ProfileAuthorityError, _verify_against_canonical


def save(document):
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def parse(document, limits=None):
    return parse_sync(save(document), "DOCX", limits or OutputLimits())


def row_document(*texts):
    document = docx.Document()
    table = document.add_table(rows=1, cols=len(texts))
    for cell, text in zip(table.rows[0].cells, texts, strict=True):
        cell.text = text
    return document


def textbox(document, text):
    # Well-understood WordprocessingML text box; never arbitrary shape/a:t text.
    run = document.add_paragraph().add_run()
    pict = OxmlElement("w:pict")
    box = OxmlElement("w:txbxContent")
    p, r, t = OxmlElement("w:p"), OxmlElement("w:r"), OxmlElement("w:t")
    t.text = text
    r.append(t)
    p.append(r)
    box.append(p)
    pict.append(box)
    run._r.append(pict)


def canonical(result):
    return SimpleNamespace(id=uuid.uuid4(), content=result.content.model_dump(mode="json"))


def claim(index, quote="Python", name="Python", category=None):
    return CandidateProfileExtraction(
        skills=[SkillItem(name=name, category=category, evidence=[ref(index, quote)])]
    )


def ref(index, quote):
    return EvidenceRef(page=1, block_index=index, quote=quote)


def test_pinned_api_order_empty_gaps_nested_multilingual_determinism():
    assert docx.__version__ == "1.2.0"
    document = docx.Document()
    document.add_paragraph("A")
    document.add_paragraph(" ")
    table = document.add_table(rows=1, cols=2)
    cell = table.cell(0, 0)
    cell.text = "Python"
    cell.add_paragraph("Azərbaycan Русский 中文")
    nested = cell.add_table(rows=1, cols=1)
    nested.cell(0, 0).text = "nested"
    cell.add_paragraph("after nested")
    table.cell(0, 1).text = "same"
    document.add_table(rows=1, cols=1).cell(0, 0).text = "same"
    document.add_paragraph("B")
    result = parse(document)
    assert result == parse(document)
    blocks = result.content.pages[0].blocks
    assert [(b.index, b.text) for b in blocks] == [
        (0, "A"),
        (2, "Python"),
        (3, "Azərbaycan Русский 中文"),
        (4, "nested"),
        (6, "after nested"),
        (7, "same"),
        (8, "same"),
        (9, "B"),
    ]
    assert all(p.page == 1 for p in result.content.pages)
    assert isinstance(blocks[0].source, BodySource)
    assert isinstance(blocks[3].source, TableSource)
    assert [step.model_dump() for step in blocks[3].source.path] == [
        {"table": 0, "row": 0, "cell": 0},
        {"table": 0, "row": 0, "cell": 0},
    ]
    assert blocks[4].source.paragraph == 3
    assert blocks[6].source.path[0].table == 1


@pytest.mark.parametrize("shape", ["horizontal", "vertical", "rectangle"])
def test_merge_origin_once_no_recursive_layout_authority(shape, monkeypatch):
    document = docx.Document()
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "origin"
    target = {"horizontal": (0, 1), "vertical": (1, 0), "rectangle": (1, 1)}[shape]
    table.cell(0, 0).merge(table.cell(*target))
    # Fail if extraction calls row.cells even for ordinary merges.
    monkeypatch.setattr(docx.table._Row, "cells", property(lambda _: pytest.fail("layout grid")))
    result = parse(document)
    assert [b.text for b in result.content.pages[0].blocks] == ["origin"]
    assert not result.content.warnings


def test_distinct_identical_cells_and_giant_gridspan_do_not_expand(monkeypatch):
    document = row_document("same", "same")
    span = document.tables[0].cell(0, 0)._tc.get_or_add_tcPr().get_or_add_gridSpan()
    span.val = 2_000_000_000
    monkeypatch.setattr(docx.table._Row, "cells", property(lambda _: pytest.fail("layout grid")))
    blocks = parse(document).content.pages[0].blocks
    assert [(b.index, b.text) for b in blocks] == [(0, "same"), (1, "same")]
    assert [b.source.path[0].cell for b in blocks] == [0, 1]


def test_long_vertical_merge_and_ambiguous_source_continuation():
    document = docx.Document()
    table = document.add_table(rows=1200, cols=1)
    for index, tr in enumerate(table._tbl.tr_lst):
        tc = tr.tc_lst[0]
        tc.get_or_add_tcPr().get_or_add_vMerge().val = "restart" if index == 0 else "continue"
    table.cell(0, 0).text = "origin"
    assert [b.text for b in parse(document).content.pages[0].blocks] == ["origin"]
    continuation = table._tbl.tr_lst[-1].tc_lst[0]
    continuation.p_lst[0].add_r().add_t("ambiguous source")
    result = parse(document)
    assert [b.text for b in result.content.pages[0].blocks] == ["origin"]
    assert result.content.warnings == ["DOCX_AMBIGUOUS_MERGE_TEXT_OMITTED"]
    view = build_professional_document_view(canonical(result))
    # Omission in the last physical row no longer invalidates the complete origin row.
    verify_extraction_evidence(view, claim(0, "origin", "origin"))


def nested_document(depth):
    document = docx.Document()
    container = document
    for _ in range(depth):
        container = container.add_table(rows=1, cols=1).cell(0, 0)
    container.text = "Python"
    return document


def test_exact_over_nesting_bound():
    result = parse(nested_document(MAX_DOCX_TABLE_DEPTH))
    assert len(result.content.pages[0].blocks[0].source.path) == MAX_DOCX_TABLE_DEPTH
    with pytest.raises(ParseError) as caught:
        parse(nested_document(MAX_DOCX_TABLE_DEPTH + 1))
    assert caught.value.code == ParseFailureCode.PARSER_OUTPUT_LIMIT


def test_structure_work_bounded_even_without_emitted_blocks(monkeypatch):
    document = row_document("")
    nodes = sum(1 for _ in document.element.iter())
    monkeypatch.setattr(docx_source, "MAX_DOCX_SOURCE_NODES", nodes)
    with pytest.raises(ParseError) as caught:
        parse(document)
    assert caught.value.code == ParseFailureCode.INSUFFICIENT_EXTRACTABLE_TEXT
    monkeypatch.setattr(docx_source, "MAX_DOCX_SOURCE_NODES", nodes - 1)
    with pytest.raises(ParseError) as caught:
        parse(document)
    assert caught.value.code == ParseFailureCode.PARSER_OUTPUT_LIMIT


@pytest.mark.parametrize("kind", ["header", "footer", "textbox"])
@pytest.mark.parametrize("body", [False, True])
def test_known_omitted_text_partial_or_unsupported_only(kind, body):
    document = docx.Document()
    if body:
        document.add_paragraph("Python")
    if kind == "textbox":
        textbox(document, "OMITTED SYNTHETIC SECRET")
    else:
        getattr(document.sections[0], kind).paragraphs[0].text = "OMITTED SYNTHETIC SECRET"
    code = f"DOCX_{kind.upper()}_TEXT_OMITTED"
    if body:
        result = parse(document)
        assert result.content.warnings == [code]
        assert [b.text for b in result.content.pages[0].blocks] == ["Python"]
        encoded = json.dumps(result.content.warnings)
        assert "SECRET" not in encoded and ".xml" not in encoded and "/" not in encoded
    else:
        with pytest.raises(ParseError) as caught:
            parse(document)
        assert caught.value.code == ParseFailureCode.UNSUPPORTED_DOCX_TEXT_ONLY
        assert caught.value.is_terminal
        assert str(caught.value) == (
            "This document contains text in DOCX structures that are not yet supported."
        )


@pytest.mark.parametrize("kind", ["header", "footer"])
def test_existing_header_variants_shared_parts_and_no_definition_creation(kind):
    document = row_document("Python")
    for variant in (kind, f"first_page_{kind}", f"even_page_{kind}"):
        getattr(document.sections[0], variant).paragraphs[0].text = "synthetic omitted"
    document.add_section()  # inherits shared definitions
    before = save(document)
    result = parse(document)
    assert result.content.warnings == [f"DOCX_{kind.upper()}_TEXT_OMITTED"]
    assert save(document) == before
    fresh = row_document("Python")
    before_parts = len(fresh.part.rels)
    docx_source.extract_body(fresh, TextBudget(OutputLimits()))
    assert len(fresh.part.rels) == before_parts


def test_blank_decorative_shapes_and_text_after_box_are_not_omitted():
    document = docx.Document()
    textbox(document, " ")
    document.add_paragraph("Python")
    result = parse(document)
    assert not result.content.warnings
    assert [b.text for b in result.content.pages[0].blocks] == ["Python"]


@pytest.mark.parametrize(
    "texts",
    [
        ("No experience with", "Python"),
        ("Python", "No experience"),
        ("Python", "Python is not available"),
    ],
)
def test_table_negative_and_contradictory_siblings_veto(texts):
    view = build_professional_document_view(canonical(parse(row_document(*texts))))
    index = next(b.block_index for b in view.blocks if b.text == "Python")
    with pytest.raises(EvidenceValidationError):
        verify_extraction_evidence(view, claim(index))
    assert "source" not in view.model_dump_json()


def test_positive_table_and_no_positive_borrowing_quote_mismatch_body_unchanged():
    view = build_professional_document_view(canonical(parse(row_document("Python", "Backend"))))
    verify_extraction_evidence(view, claim(0))
    for extraction in (claim(0, category="Backend"), claim(1), claim(0, "Java")):
        with pytest.raises(EvidenceValidationError):
            verify_extraction_evidence(view, extraction)
    # Employment material values must all remain in its own cited block.
    with pytest.raises(EvidenceValidationError):
        verify_extraction_evidence(
            view,
            CandidateProfileExtraction(
                employment_history=[
                    EmploymentItem(
                        title="Backend", organization="Python", evidence=[ref(1, "Backend")]
                    )
                ]
            ),
        )
    body = docx.Document()
    body.add_paragraph("No experience with")
    body.add_paragraph("Python")
    verify_extraction_evidence(build_professional_document_view(canonical(parse(body))), claim(1))


@pytest.mark.parametrize("label", ["employee ID", "reference", "account", "code"])
def test_identity_table_nonphone_label_veto(label):
    view = build_identity_document_view(canonical(parse(row_document(label, "+994 50 123 45 67"))))
    extraction = CandidateIdentityExtraction(
        phone=IdentityFieldItem(value="+994 50 123 45 67", evidence=[ref(1, "+994 50 123 45 67")])
    )
    with pytest.raises(EvidenceValidationError):
        verify_identity_evidence(view, extraction)


def test_positive_phone_label_cannot_grant_sibling_numeric_authority():
    view = build_identity_document_view(canonical(parse(row_document("Phone", "994501234567"))))
    with pytest.raises(EvidenceValidationError):
        verify_identity_evidence(
            view,
            CandidateIdentityExtraction(
                phone=IdentityFieldItem(value="994501234567", evidence=[ref(1, "994501234567")])
            ),
        )
    view = build_identity_document_view(canonical(parse(row_document("Email", "test@example.com"))))
    verify_identity_evidence(
        view,
        CandidateIdentityExtraction(
            email=IdentityFieldItem(value="test@example.com", evidence=[ref(1, "test@example.com")])
        ),
    )


def test_current_authority_rebuild_uses_table_veto_without_mutating_rows():
    source = canonical(parse(row_document("No experience with", "Python")))
    source.tenant_id = uuid.uuid4()
    source.candidate_document_id = uuid.uuid4()
    version = SimpleNamespace(
        canonical_document_id=source.id,
        tenant_id=source.tenant_id,
        candidate_document_id=source.candidate_document_id,
    )
    before = json.dumps(source.content)
    with pytest.raises(ProfileAuthorityError):
        _verify_against_canonical(version, claim(1), source)
    assert json.dumps(source.content) == before


@pytest.mark.parametrize("field,exact", [("blocks", 2), ("characters", 12), ("text_bytes", 10)])
def test_table_text_budgets_exact_over(field, exact):
    texts = ("Python", "Python") if field != "text_bytes" else ("中文", "əə")
    limits = replace(OutputLimits(), **{field: exact})
    result = parse(row_document(*texts), limits)
    with pytest.raises(ParseError) as caught:
        parse(row_document(*texts), replace(limits, **{field: exact - 1}))
    assert caught.value.code == ParseFailureCode.PARSER_OUTPUT_LIMIT
    assert result.content.pages[0].blocks


def test_serialized_budget_includes_provenance_and_warnings():
    document = row_document("Python")
    document.sections[0].header.paragraphs[0].text = "omitted"
    result = parse(document)
    encoded = parser_worker.encode_bounded(result.model_dump(mode="json"), 10000)
    parser_supervisor.decode_result(
        encoded, "DOCX", replace(OutputLimits(), result_bytes=len(encoded))
    )
    for operation in (
        lambda: parser_worker.encode_bounded(result.model_dump(mode="json"), len(encoded) - 1),
        lambda: parser_supervisor.decode_result(
            encoded, "DOCX", replace(OutputLimits(), result_bytes=len(encoded) - 1)
        ),
    ):
        with pytest.raises(ParseError) as caught:
            operation()
        assert caught.value.code == ParseFailureCode.PARSER_OUTPUT_LIMIT


@pytest.mark.parametrize(
    "fault",
    [
        "unknown",
        "extra",
        "negative",
        "bool",
        "string",
        "empty",
        "deep",
        "duplicate",
        "warning_unknown",
        "warning_duplicate",
        "warnings_over",
    ],
)
def test_parent_strict_provenance_and_warnings(fault):
    payload = parse(row_document("Python")).model_dump(mode="json")
    content = payload["content"]
    block = content["pages"][0]["blocks"][0]
    source = block["source"]
    if fault == "unknown":
        source["kind"] = "XML"
    if fault == "extra":
        source["xml_path"] = "secret"
    if fault in ("negative", "bool", "string"):
        source["path"][0]["row"] = {"negative": -1, "bool": True, "string": "0"}[fault]
    if fault == "empty":
        source["path"] = []
    if fault == "deep":
        source["path"] *= MAX_DOCX_TABLE_DEPTH + 1
    if fault == "duplicate":
        content["pages"][0]["blocks"].append({**block, "index": 1})
    if fault == "warning_unknown":
        content["warnings"] = ["PRIVATE_XML"]
    if fault == "warning_duplicate":
        content["warnings"] = ["DOCX_HEADER_TEXT_OMITTED"] * 2
    if fault == "warnings_over":
        content["warnings"] = ["DOCX_HEADER_TEXT_OMITTED"] * 6
    with pytest.raises(ParseError) as caught:
        parser_supervisor.decode_result(json.dumps(payload).encode(), "DOCX", OutputLimits())
    assert caught.value.code == ParseFailureCode.INVALID_PARSER_OUTPUT
    if fault == "warnings_over":
        with pytest.raises(ValidationError):
            type(parse(row_document("Python")).content).model_validate(content)


def test_table_veto_context_exact_over_and_other_rows_not_borrowed():
    from meyar.extraction import evidence

    for count in (evidence._TABLE_CONTEXT_BLOCKS, evidence._TABLE_CONTEXT_BLOCKS + 1):
        view = build_professional_document_view(
            canonical(parse(row_document("Python", *(["label"] * (count - 1)))))
        )
        if count == evidence._TABLE_CONTEXT_BLOCKS:
            verify_extraction_evidence(view, claim(0))
        else:
            with pytest.raises(EvidenceValidationError):
                verify_extraction_evidence(view, claim(0))
    for total in (evidence._TABLE_CONTEXT_CHARS, evidence._TABLE_CONTEXT_CHARS + 1):
        view = build_professional_document_view(
            canonical(parse(row_document("Python", "x" * (total - 6))))
        )
        if total == evidence._TABLE_CONTEXT_CHARS:
            verify_extraction_evidence(view, claim(0))
        else:
            with pytest.raises(EvidenceValidationError):
                verify_extraction_evidence(view, claim(0))
    document = row_document("Python")
    document.tables[0].add_row().cells[0].text = "No experience with Python"
    verify_extraction_evidence(
        build_professional_document_view(canonical(parse(document))), claim(0)
    )


def test_actual_source_node_policy_exact_over_empty_work():
    from meyar.ingestion.parser_policy import MAX_DOCX_SOURCE_NODES

    root = OxmlElement("w:body")
    for _ in range(MAX_DOCX_SOURCE_NODES - 1):
        root.append(OxmlElement("w:p"))
    docx_source.SourceBudget().inspect(root, set())
    root.append(OxmlElement("w:p"))
    with pytest.raises(ParseError) as caught:
        docx_source.SourceBudget().inspect(root, set())
    assert caught.value.code == ParseFailureCode.PARSER_OUTPUT_LIMIT


@pytest.mark.parametrize("field", ["blocks", "characters"])
def test_actual_table_text_policy_exact_over(field):
    document = row_document("Hi" if field == "blocks" else "x" * 1_000_000)
    if field == "blocks":
        cell = document.tables[0].cell(0, 0)
        for _ in range(9999):
            cell.add_paragraph("Hi")
    assert parse(document)
    if field == "blocks":
        cell.add_paragraph("Hi")
    else:
        document.tables[0].cell(0, 0).text += "x"
    with pytest.raises(ParseError) as caught:
        parse(document)
    assert caught.value.code == ParseFailureCode.PARSER_OUTPUT_LIMIT


def test_ambiguous_continuation_only_terminal():
    document = row_document("ambiguous source")
    document.tables[0].cell(0, 0)._tc.get_or_add_tcPr().get_or_add_vMerge().val = "continue"
    with pytest.raises(ParseError) as caught:
        parse(document)
    assert caught.value.code == ParseFailureCode.UNSUPPORTED_DOCX_TEXT_ONLY


def test_continuation_unsupported_wrapper_text_detected_without_extraction():
    document = row_document("")
    tc = document.tables[0].cell(0, 0)._tc
    tc.get_or_add_tcPr().get_or_add_vMerge().val = "continue"
    wrapper = OxmlElement("w:sdt")
    content = OxmlElement("w:sdtContent")
    p, r, t = OxmlElement("w:p"), OxmlElement("w:r"), OxmlElement("w:t")
    t.text = "synthetic omitted wrapper"
    r.append(t)
    p.append(r)
    content.append(p)
    wrapper.append(content)
    tc.append(wrapper)
    with pytest.raises(ParseError) as caught:
        parse(document)
    assert caught.value.code == ParseFailureCode.UNSUPPORTED_DOCX_TEXT_ONLY
    document.add_paragraph("Python")
    result = parse(document)
    assert result.content.warnings == ["DOCX_AMBIGUOUS_MERGE_TEXT_OMITTED"]
    assert [b.text for b in result.content.pages[0].blocks] == ["Python"]


def test_all_five_closed_warnings_exact_cardinality():
    from test_docx_review_corrections import omitted

    document = row_document("ambiguous source")
    document.tables[0].cell(0, 0)._tc.get_or_add_tcPr().get_or_add_vMerge().val = "continue"
    document.add_paragraph("Python")
    document.sections[0].header.paragraphs[0].text = "omitted header"
    document.sections[0].footer.paragraphs[0].text = "omitted footer"
    textbox(document, "omitted textbox")
    omitted(document.add_table(rows=1, cols=1).cell(0, 0), "sdt")
    result = parse(document)
    assert result.content.warnings == [
        "DOCX_HEADER_TEXT_OMITTED",
        "DOCX_FOOTER_TEXT_OMITTED",
        "DOCX_TEXTBOX_TEXT_OMITTED",
        "DOCX_AMBIGUOUS_MERGE_TEXT_OMITTED",
        "DOCX_TABLE_TEXT_OMITTED",
    ]
    assert parser_supervisor.decode_result(
        result.model_dump_json().encode(), "DOCX", OutputLimits()
    )


def test_unsupported_docx_only_error_cannot_be_pdf_terminal_authority():
    with pytest.raises(ParseError) as caught:
        parser_supervisor.decode_result(
            b'{"error":"UNSUPPORTED_DOCX_TEXT_ONLY"}', "PDF", OutputLimits()
        )
    assert caught.value.code == ParseFailureCode.INVALID_PARSER_OUTPUT


def test_redaction_cannot_erase_server_only_table_negative_context():
    view = build_professional_document_view(
        canonical(parse(row_document("DOB: No experience with", "Python")))
    )
    assert view.blocks[0].text == "[REDACTED_FIELD]"
    assert "DOB" not in view.model_dump_json()
    assert "table_context_text" not in view.model_dump_json()
    with pytest.raises(EvidenceValidationError):
        verify_extraction_evidence(view, claim(1))
